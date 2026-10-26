import io
from openpyxl import load_workbook
from fastapi import UploadFile
from pydantic import ValidationError
from app.schemas.vocabulary import VocabularyImportRow


class VocabExcelImporter:

    @staticmethod
    async def parse_excel(file: UploadFile) -> tuple[list[VocabularyImportRow], list[dict]]:
        content = await file.read()
        
        try:
            # data_only=True để lấy giá trị thực tế (nếu ô có công thức hàm)
            workbook = load_workbook(filename=io.BytesIO(content), data_only=True)
            sheet = workbook.worksheets[0] 
        except Exception:
            raise ValueError("Không thể đọc file Excel. Vui lòng đảm bảo file định dạng .xlsx hợp lệ.")

        # Lấy tất cả các dòng dưới dạng giá trị
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            raise ValueError("File Excel trống hoặc không có dòng dữ liệu nào.")

        # Dòng 1 là tiêu đề cột (Header)
        headers = [str(col).strip() for col in rows[0] if col is not None]
        
        valid_rows: list[VocabularyImportRow] = []
        errors: list[dict] = []

        # Bắt đầu duyệt từ dòng 2 (index 1 trong mảng, trên Excel là dòng 2)
        for row_idx, row in enumerate(rows[1:], start=2):
            # Bỏ qua nếu dòng hoàn toàn trống
            if not any(row):
                continue

            # Ghép header với giá trị ô tương ứng
            row_dict = {}
            for col_idx, header in enumerate(headers):
                if col_idx < len(row):
                    val = row[col_idx]
                    # Chuyển đổi giá trị về dạng chuỗi sạch
                    row_dict[header] = str(val).strip() if val is not None else ""

            # Validate bằng Pydantic
            try:
                item = VocabularyImportRow(**row_dict)
                valid_rows.append(item)
            except ValidationError as e:
                error_msg = "; ".join([f"{err['loc'][0]}: {err['msg']}" for err in e.errors()])
                errors.append({
                    "row": row_idx,
                    "word": row_dict.get("word", "N/A"),
                    "error": error_msg
                })

        return valid_rows, errors
    
    